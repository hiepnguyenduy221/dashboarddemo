"""Đọc Excel theo template dashboard và ghi data.json.

    python tools/excel_to_json.py duong_dan.xlsx            # ghi ra data.json ở gốc repo
    python tools/excel_to_json.py duong_dan.xlsx -o out.json
    python tools/excel_to_json.py duong_dan.xlsx --check    # chỉ kiểm tra, không ghi

Có lỗi → không ghi file, thoát mã 1. Cảnh báo (đối soát lệch) vẫn ghi file.
Cần: pip install openpyxl
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Thiếu thư viện openpyxl: pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dashboard_schema import (  # noqa: E402
    COMPANY, META_KEYS, SHEETS, SHEET_BY_NAME, column_codes,
)

ROOT = Path(__file__).resolve().parent.parent
FORMAT_VERSION = 1


class Report:
    def __init__(self):
        self.errors, self.warnings = [], []

    def error(self, where, msg):
        self.errors.append(f"[{where}] {msg}")

    def warn(self, where, msg):
        self.warnings.append(f"[{where}] {msg}")


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"ngày không hợp lệ '{text}' (dùng dd/mm/yyyy)")


def parse_number(value, integer):
    if isinstance(value, bool):
        raise ValueError("không phải số")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(" ", "").replace(" ", "")
        if text.endswith("%"):
            text = text[:-1]
        # Chấp nhận cả 1.234,56 (VN) và 1234.56
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            raise ValueError(f"'{value}' không phải số") from None
    if integer:
        if abs(number - round(number)) > 1e-9:
            raise ValueError(f"cần số nguyên, nhận {value}")
        return int(round(number))
    return round(number, 6)


def parse_cell(value, kind, units):
    """Trả về giá trị đã chuẩn hóa; ô trống → None."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if kind == "any":
        return value.strip() if isinstance(value, str) else value
    if kind == "date":
        return parse_date(value).isoformat()
    if kind in ("int", "num"):
        return parse_number(value, kind == "int")
    if kind == "flag":
        text = str(value).strip().lower()
        if text in ("x", "1", "true", "có", "co", "yes"):
            return True
        if text in ("0", "false", "không", "khong", "no"):
            return None
        raise ValueError(f"cờ chỉ nhận 'x', nhận '{value}'")
    text = str(value).strip()
    if kind.startswith("enum:"):
        allowed = kind[5:].split("|")
        if text not in allowed:
            raise ValueError(f"'{text}' không thuộc {', '.join(allowed)}")
    elif kind.startswith("unit"):
        allowed = set(units)
        if kind == "unit+company":
            allowed.add(COMPANY)
        if text not in allowed:
            raise ValueError(f"đơn vị '{text}' chưa có trong DON_VI")
    return text


def read_sheet(wb, sheet, units, report):
    name = sheet["name"]
    if name not in wb.sheetnames:
        report.error(name, "thiếu sheet")
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        report.error(name, "sheet trống")
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    specs = {c[0]: c for c in sheet["columns"]}
    index = {}
    for i, code in enumerate(header):
        if not code:
            continue
        if code in index:
            report.error(name, f"cột '{code}' bị lặp")
        elif code not in specs:
            report.warn(name, f"bỏ qua cột lạ '{code}'")
        else:
            index[code] = i
    for code, (_, _, _, required, _) in specs.items():
        if required and code not in index:
            report.error(name, f"thiếu cột bắt buộc '{code}'")
    if report.errors and any(e.startswith(f"[{name}] thiếu cột") for e in report.errors):
        return []

    body = rows[1:]
    first_desc = sheet["columns"][0][4]
    if body and body[0] and str(body[0][0] or "").strip() == first_desc:
        body = body[1:]  # dòng 2 là mô tả cột
        first_line = 3
    else:
        first_line = 2

    out, seen = [], {}
    for offset, raw in enumerate(body):
        line = first_line + offset
        if raw is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
            continue
        record, bad = {}, False
        for code, i in index.items():
            kind, required = specs[code][1], specs[code][3]
            value = raw[i] if i < len(raw) else None
            try:
                parsed = parse_cell(value, kind, units)
            except (ValueError, TypeError) as exc:
                report.error(name, f"dòng {line}, cột {code}: {exc}")
                bad = True
                continue
            if parsed is None:
                if required:
                    report.error(name, f"dòng {line}: thiếu '{code}'")
                    bad = True
                continue
            record[code] = parsed
        if bad:
            continue
        key = tuple(record.get(k, "") for k in sheet["key"])
        if key in seen:
            report.error(name, f"dòng {line} trùng khóa {dict(zip(sheet['key'], key))} với dòng {seen[key]}")
            continue
        seen[key] = line
        out.append(record)
    return out


def read_meta(rows, report):
    meta = {}
    known = {k[0]: k for k in META_KEYS}
    for row in rows:
        key = row["khoa"]
        if key not in known:
            report.warn("THONG_TIN", f"bỏ qua khóa lạ '{key}'")
            continue
        kind = known[key][1]
        try:
            meta[key] = parse_cell(row.get("gia_tri"), kind, [])
        except (ValueError, TypeError) as exc:
            report.error("THONG_TIN", f"{key}: {exc}")
    for key, _, required, _ in META_KEYS:
        if required and meta.get(key) is None:
            report.error("THONG_TIN", f"thiếu giá trị '{key}'")
    return meta


def close(a, b, tolerance):
    return a is None or b is None or abs(a - b) <= tolerance


def cross_check(data, report):
    meta = data["thong_tin"]
    report_date = meta.get("ngay_bao_cao")
    unit_names = [u["ten_don_vi"] for u in data["don_vi"]]

    days = data["ngay"]
    if report_date:
        month = report_date[:7]
        for row in days:
            if row["ngay"] > report_date or row["ngay"][:7] != month:
                report.error("NGAY", f"phiên {row['ngay']} nằm ngoài tháng báo cáo {month}")
        if not any(r["ngay"] == report_date and r["don_vi"] == COMPANY for r in days):
            report.error("NGAY", f"thiếu dòng 'Toàn công ty' ngày {report_date}")
    dates = sorted({r["ngay"] for r in days})
    for unit in [COMPANY] + unit_names:
        have = {r["ngay"] for r in days if r["don_vi"] == unit}
        missing = [d for d in dates if d not in have]
        if missing and len(missing) < len(dates):
            report.warn("NGAY", f"{unit}: thiếu {len(missing)} phiên ({', '.join(missing[:3])}…)")
        if not have:
            report.warn("NGAY", f"{unit}: chưa có dòng nào")

    ky = {r["don_vi"]: r for r in data["ky"]}
    if COMPANY not in ky:
        report.error("KY", "thiếu dòng 'Toàn công ty'")
    today = {r["don_vi"]: r for r in days if r["ngay"] == report_date}
    for unit, row in ky.items():
        day = today.get(unit, {})
        parts = [row.get(k) for k in ("du_no_margin", "du_no_ung_truoc", "du_no_3_ben")]
        if all(p is not None for p in parts) and not close(sum(parts), day.get("du_no"), 1.5):
            report.warn("KY", f"{unit}: margin + ứng trước + 3 bên = {sum(parts):,.2f} ≠ dư nợ ngày {day.get('du_no')}")
        if row.get("tien_nap") is not None and row.get("tien_rut") is not None:
            net = row["tien_nap"] - row["tien_rut"]
            if not close(net, day.get("nap_rut_rong"), 1):
                report.warn("KY", f"{unit}: nạp − rút = {net:,.2f} ≠ nạp/rút ròng ngày {day.get('nap_rut_rong')}")

    kh_units = {r["don_vi"] for r in data["kh"]}
    for unit in unit_names:
        if unit not in kh_units:
            report.warn("KH", f"{unit}: chưa có dòng")

    for name, key in (("ABCD", "abcd"), ("PKKH", "pkkh")):
        if not data[key]:
            report.warn(name, "sheet chưa có dữ liệu — phần này sẽ hiển thị '—'")

    staff_total = {}
    for row in data["phong"]:
        unit = row.get("don_vi", "")
        fee = sum(row.get(k) or 0 for k in ("dtpn_mg", "dtpn_ctv", "dtpn_chua_phan_loai"))
        staff_total[unit] = staff_total.get(unit, 0) + fee
    if staff_total and meta.get("ngay_chot_nhan_su") == report_date:
        for unit, fee in staff_total.items():
            if not unit:
                continue
            business = sum(r.get("phi_net") or 0 for r in days if r["don_vi"] == unit)
            if business and abs(fee - business) > max(5, business * 0.02):
                report.warn("PHONG", f"{unit}: tổng DTPN phòng {fee:,.1f} triệu lệch phí lũy kế {business:,.1f} triệu")


def convert(path, report):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_rows = {}
    units = []
    for name in ("THONG_TIN", "DON_VI"):
        sheet_rows[name] = read_sheet(wb, SHEET_BY_NAME[name], units, report)
    units = [r["ten_don_vi"] for r in sheet_rows["DON_VI"]]
    if COMPANY in units:
        report.error("DON_VI", f"không khai báo '{COMPANY}' như một đơn vị")
    for sheet in SHEETS:
        if sheet["name"] not in sheet_rows:
            sheet_rows[sheet["name"]] = read_sheet(wb, sheet, units, report)

    units_rows = sorted(
        enumerate(sheet_rows["DON_VI"]),
        key=lambda item: (item[1].get("thu_tu", 10**6), item[0]),
    )
    data = {
        "phien_ban": FORMAT_VERSION,
        "tao_luc": dt.datetime.now().isoformat(timespec="seconds"),
        "tu_file": Path(path).name,
        "thong_tin": read_meta(sheet_rows["THONG_TIN"], report),
        "don_vi": [{k: v for k, v in row.items() if k != "thu_tu"} for _, row in units_rows],
    }
    for sheet in SHEETS[2:]:
        # Giữ thứ tự dòng trong Excel: thứ tự phòng/đơn vị trên dashboard theo đúng file.
        data[sheet["name"].lower()] = sheet_rows[sheet["name"]]
    cross_check(data, report)
    return data


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", help="File Excel theo template")
    parser.add_argument("-o", "--output", default=str(ROOT / "data.json"), help="File JSON đầu ra (mặc định data.json)")
    parser.add_argument("--check", action="store_true", help="Chỉ kiểm tra, không ghi file")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    report = Report()
    data = convert(args.excel, report)
    for w in report.warnings:
        print("CẢNH BÁO", w)
    for e in report.errors:
        print("LỖI", e)
    if report.errors:
        print(f"\n{len(report.errors)} lỗi — không ghi {args.output}.")
        return 1
    counts = ", ".join(f"{s['name']} {len(data[s['name'].lower()])}" for s in SHEETS[2:])
    if args.check:
        print(f"\nHợp lệ ({counts}).")
        return 0
    Path(args.output).write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"\nĐã ghi {args.output} — ngày báo cáo {data['thong_tin'].get('ngay_bao_cao')}; {counts}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
