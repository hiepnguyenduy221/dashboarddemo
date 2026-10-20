"""Sinh file Excel template cho team data.

    python tools/make_template.py                          # data/template_dashboard.xlsx (trống)
    python tools/make_template.py --from-json data.json -o data/mau.xlsx   # template đã điền từ data.json

Cần: pip install openpyxl
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("Thiếu thư viện openpyxl: pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dashboard_schema import COMPANY, META_KEYS, SHEETS  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
NAVY = PatternFill("solid", fgColor="102C50")
SOFT = PatternFill("solid", fgColor="EEF2F7")
REQ = PatternFill("solid", fgColor="F26A21")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
DESC_FONT = Font(italic=True, color="64758B", size=9)
MAX_ROWS = 3000
NUMBER_FORMATS = {"int": "#,##0", "num": "#,##0.00", "date": "dd/mm/yyyy"}


def write_guide(wb):
    ws = wb.active
    ws.title = "HUONG_DAN"
    lines = [
        ("TEMPLATE DỮ LIỆU DASHBOARD CEO", True),
        ("", False),
        ("Cách dùng", True),
        ("1. Điền THONG_TIN và DON_VI trước; tên đơn vị ở DON_VI là tên duy nhất dùng ở mọi sheet.", False),
        ("2. Mỗi sheet: dòng 1 là mã cột (không sửa), dòng 2 là mô tả, dữ liệu bắt đầu từ dòng 3.", False),
        ("3. Cột tiêu đề màu cam là bắt buộc. Ô không có số liệu thì để trống — dashboard hiện '—', không hiểu là 0.", False),
        ("4. Đơn vị tính ghi trong mô tả cột: GTGD/NAV/dư nợ = tỷ đồng; phí/DTPN = triệu đồng; thị phần = % (2,41 nghĩa là 2,41%).", False),
        ("5. Không nhập số tính lại được: lũy kế tháng, AUM, Δ phiên trước, bình quân, tổng toàn công ty, tỷ lệ % — dashboard tự tính.", False),
        ("6. Chạy: python tools/excel_to_json.py <file>.xlsx  → kiểm tra và ghi data.json.", False),
        ("", False),
        ("Các sheet", True),
    ]
    for sheet in SHEETS:
        lines.append((f"{sheet['name']} — {sheet['title']}: {sheet['grain']}", False))
    lines += [("", False), (f"Đơn vị đặc biệt '{COMPANY}' chỉ dùng ở NGAY và KY (số công bố toàn công ty, không phải tổng các đơn vị).", False)]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=13 if i == 1 else 11, color="102C50" if bold else "172C43")
    ws.column_dimensions["A"].width = 140


def write_lists(wb):
    ws = wb.create_sheet("_DS")
    ws.sheet_state = "hidden"
    ws["A1"] = COMPANY
    for r in range(2, 202):
        ws.cell(row=r, column=1, value=f'=IF(DON_VI!A{r + 1}="","",DON_VI!A{r + 1})')


def write_sheet(wb, sheet, rows):
    ws = wb.create_sheet(sheet["name"])
    for col, (code, kind, unit, required, desc) in enumerate(sheet["columns"], start=1):
        head = ws.cell(row=1, column=col, value=code)
        head.font = WHITE_BOLD
        head.fill = REQ if required else NAVY
        label = desc + (f" [{unit}]" if unit else "")
        cell = ws.cell(row=2, column=col, value=desc if col == 1 else label)
        cell.font = DESC_FONT
        cell.fill = SOFT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(code) + 4))
        letter = get_column_letter(col)
        target = f"{letter}3:{letter}{MAX_ROWS}"
        formula = None
        if kind.startswith("enum:"):
            formula = '"' + ",".join(kind[5:].split("|")) + '"'
        elif kind in ("unit", "unit+blank"):
            formula = "=DON_VI!$A$3:$A$202"
        elif kind == "unit+company":
            formula = "=_DS!$A$1:$A$201"
        elif kind == "flag":
            formula = '"x"'
        if formula:
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
            ws.add_data_validation(dv)
            dv.add(target)
        for r in range(3, 3 + max(len(rows), 1)):
            if kind in NUMBER_FORMATS:
                ws.cell(row=r, column=col).number_format = NUMBER_FORMATS[kind]
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "A3"

    specs = sheet["columns"]
    for r, row in enumerate(rows, start=3):
        for col, (code, kind, *_rest) in enumerate(specs, start=1):
            value = row.get(code)
            if value is None:
                continue
            if kind == "date" or (sheet["name"] == "THONG_TIN" and isinstance(value, str) and _looks_iso(value)):
                value = dt.date.fromisoformat(value)
                ws.cell(row=r, column=col).number_format = "dd/mm/yyyy"
            elif kind == "flag":
                value = "x" if value else None
            ws.cell(row=r, column=col, value=value)


def _looks_iso(text):
    try:
        dt.date.fromisoformat(text)
        return len(text) == 10
    except ValueError:
        return False


def build(data, output):
    wb = openpyxl.Workbook()
    write_guide(wb)
    meta = (data or {}).get("thong_tin", {})
    rows_by_sheet = {
        "THONG_TIN": [{"khoa": k, "gia_tri": meta.get(k), "mo_ta": desc} for k, _, _, desc in META_KEYS],
    }
    if data:
        rows_by_sheet["DON_VI"] = [dict(r, thu_tu=i + 1) for i, r in enumerate(data["don_vi"])]
        for sheet in SHEETS[2:]:
            rows_by_sheet[sheet["name"]] = data.get(sheet["name"].lower(), [])
    for sheet in SHEETS:
        write_sheet(wb, sheet, rows_by_sheet.get(sheet["name"], []))
    write_lists(wb)
    wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


SCHEMA_OPEN = '<script id="schema" type="application/json">'


def sync_import_page():
    """Chép schema vào tools/import.html để trang import kiểm tra đúng như script Python."""
    page = Path(__file__).resolve().parent / "import.html"
    if not page.exists():
        return
    schema = {
        "company": COMPANY,
        "meta_keys": [list(k) for k in META_KEYS],
        "sheets": [{"name": s["name"], "key": s["key"], "columns": [list(c) for c in s["columns"]]} for s in SHEETS],
    }
    html = page.read_text(encoding="utf-8")
    start = html.index(SCHEMA_OPEN) + len(SCHEMA_OPEN)
    end = html.index("</script>", start)
    payload = json.dumps(schema, ensure_ascii=False).replace("</", "<\\/")
    page.write_text(html[:start] + payload + html[end:], encoding="utf-8", newline="\n")
    print(f"Đã cập nhật schema trong {page}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--from-json", help="Điền sẵn từ data.json")
    parser.add_argument("-o", "--output", default=str(ROOT / "data" / "template_dashboard.xlsx"))
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    data = json.loads(Path(args.from_json).read_text(encoding="utf-8")) if args.from_json else None
    build(data, Path(args.output))
    print(f"Đã tạo {args.output}")
    if not args.from_json:
        sync_import_page()


if __name__ == "__main__":
    main()
